from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from app.engine.analyzer import ThreatAnalyzer
from app.engine.parser import ArchitectureParser
from app.services import document_ingestion


ROOT = Path(__file__).resolve().parents[2]
COMPLEX_SCENARIO = ROOT / "test-documents" / "Complex_Healthcare_AI_Threat_Model_Scenario.docx"

# The docx these checks were written against is not in the repository, so they
# only run where someone has restored it. tests/fixtures/reference_architecture.txt
# is a text reconstruction of the same system carrying the same record counts,
# and test_authoritative_id_independence.py exercises the parser against it on
# every run; these remain as the check that a real docx still ingests correctly.
requires_complex_scenario = pytest.mark.skipif(
    not COMPLEX_SCENARIO.exists(),
    reason=f"missing regression document: {COMPLEX_SCENARIO.name}",
)


def _extracted_scenario():
    text, metadata = document_ingestion._extract_docx_text(COMPLEX_SCENARIO.read_bytes())
    description = (
        f"Document: {COMPLEX_SCENARIO.name}\n"
        "Type: docx\nRole: source_design\nContent:\n"
        f"{text}"
    )
    return description, metadata


@requires_complex_scenario
def test_docx_extraction_preserves_table_order_and_structure():
    description, metadata = _extracted_scenario()

    assert metadata["extraction_quality"] == "structured_text_complete"
    assert metadata["tables"] == "15"
    assert "[Table 5]" in description
    assert "C25 | Delivery platform" in description
    assert "F20 | Support engineer -> support portal -> tenant data" in description
    assert "K30 | Privacy / deletion" in description
    assert description.index("3. Architecture inventory") < description.index("[Table 5]")
    assert description.index("8. Known weaknesses") < description.index("K1 | Authorization")


def test_docx_fallback_preserves_tables_without_python_docx(monkeypatch):
    document = Document()
    document.add_paragraph("3. Architecture inventory")
    table = document.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text = "ID"
    table.rows[0].cells[1].text = "Component"
    table.rows[0].cells[2].text = "Technology"
    table.rows[1].cells[0].text = "C1"
    table.rows[1].cells[1].text = "Core API"
    table.rows[1].cells[2].text = "Node.js"
    payload = BytesIO()
    document.save(payload)

    monkeypatch.setattr(document_ingestion, "Document", None)
    text, metadata = document_ingestion._extract_docx_text(payload.getvalue())

    assert metadata["extraction_quality"] == "structured_text_fallback"
    assert metadata["tables"] == "1"
    assert text.index("3. Architecture inventory") < text.index("[Table 1]")
    assert "Row 2: C1 | Core API | Node.js" in text


def test_architecture_yaml_preserves_ids_flows_and_known_issues():
    payload = b"""
system: Claims SaaS
components:
  - id: web
    name: Customer Portal
    type: WebClient
    trust: public
  - id: api
    name: Claims API
    type: API
    trust: internal
  - id: database
    name: Claims Database
    type: Database
    trust: restricted
flows:
  - source: web
    target: api
    protocol: HTTPS
    data: access token and claim request
  - source: api
    target: database
    protocol: TLS
    data: customer PII
known_issues:
  - Claims are loaded by caller-provided tenant_id without server-side ownership validation.
"""
    text, metadata = document_ingestion._extract_structured_text(payload, ".yml")
    architecture = ArchitectureParser().parse(
        "Document: architecture.yml\nType: yml\nRole: source_design\nContent:\n" + text
    )

    assert metadata["structured_kind"] == "architecture"
    assert {component.id for component in architecture.components} == {"web", "api", "database"}
    assert {(flow.source_id, flow.target_id) for flow in architecture.flows} == {
        ("web", "api"), ("api", "database"),
    }
    assert all(not flow.assumed for flow in architecture.flows)
    assert architecture.metadata["known_issues"][0]["suggested_threat_id"] == \
        "API-BOLA-TENANT-CONTROL-001"


@requires_complex_scenario
def test_authoritative_tables_replace_heuristic_topology():
    description, _ = _extracted_scenario()
    architecture = ArchitectureParser().parse(description)

    technical = [component for component in architecture.components if not component.properties.get("external")]
    external = [component for component in architecture.components if component.properties.get("external")]
    flow_ids = {flow.properties.get("source_record_id") for flow in architecture.flows}
    component_ids = {component.id for component in architecture.components}

    assert len(technical) == 25
    assert {"Stripe API", "SendGrid", "Insurer lab"}.issubset(
        {component.name for component in external}
    )
    assert len(architecture.flows) == 20
    assert flow_ids == {f"F{index}" for index in range(1, 21)}
    assert all(not flow.assumed for flow in architecture.flows)
    assert len(architecture.trust_boundaries) == 9
    assert len(architecture.assets) == 12
    assert len(architecture.metadata["actors"]) == 10
    assert len(architecture.metadata["known_issues"]) == 30
    assert "ups_external" not in component_ids
    assert "vault" not in component_ids
    by_source_id = {item.properties.get("source_record_id"): item for item in technical}
    assert by_source_id["C16"].trust_level == "restricted"
    assert by_source_id["C15"].properties["public_access"] is False


@requires_complex_scenario
def test_complex_scenario_findings_are_complete_and_grounded():
    description, metadata = _extracted_scenario()
    result = ThreatAnalyzer().analyze_from_text(
        description,
        project_name="Complex document regression",
        use_local_slm=False,
        analysis_mode="standard",
        domain_profile="healthcare",
        source_documents=[{
            "filename": COMPLEX_SCENARIO.name,
            "type": "docx",
            "role": "source_design",
            **metadata,
        }],
    )

    confirmed = [threat for threat in result.threats if threat.tier == "Confirmed"]
    potential = [threat for threat in result.threats if threat.tier == "Potential"]
    assert len(confirmed) == 30
    assert len(potential) >= 6
    assert all(any(detail.get("source_ref", "").startswith("K") for detail in threat.evidence_details)
               for threat in confirmed)
    assert {threat.stride_category for threat in result.threats} == {
        "Spoofing", "Tampering", "Repudiation", "Information Disclosure",
        "Denial of Service", "Elevation of Privilege",
    }
    assert not any(threat.id.startswith("KB-K8S-005") for threat in result.threats)
    # An explicit weakness does not answer independent authentication and token
    # lifecycle questions on the same component.
    assert all(threat.id.startswith("STRIDE-") or threat.id in {
        "CTX-OAUTH-001", "CTX-FHIR-001", "CTX-SESSION-001",
    } for threat in potential)
    assert not any("public access" in threat.title.lower() and threat.affected_component == "c15"
                   for threat in result.threats)

    known_ids = {threat.id.rsplit("-K", 1)[0] for threat in confirmed}
    assert {
        "API-BOLA-TENANT-CONTROL-001",
        "AUTH-SESSION-REVOCATION-001",
        "FHIR-PARTNER-SPOOFING-001",
        "WEB-SQL-INJECTION-ORDER-001",
        "WEB-STORED-XSS-001",
        "WEB-SSRF-URL-FETCH-001",
        "AI-RAG-TENANT-ISOLATION-001",
        "AI-INDIRECT-PROMPT-INJECTION-001",
        "MCP-DELEGATED-AUTHORIZATION-001",
        "PAYMENT-IDEMPOTENCY-001",
        "SUPPLY-CHAIN-GITHUB-OIDC-001",
        "DATA-DELETION-PROPAGATION-001",
    }.issubset(known_ids)

    assert result.architecture_validation["counts"] == {
        "components": 34,
        "explicit_components": 34,
        "flows": 20,
        "explicit_flows": 20,
        "inferred_flows": 0,
        "actors": 10,
        "identities": 10,
    }
    assert "React patient web app" in result.mermaid_diagram
    assert "Public API edge" in result.mermaid_diagram
    assert "Core API" in result.mermaid_diagram
    assert "Transactional database" in result.mermaid_diagram
    assert "AI orchestrator" in result.mermaid_diagram
