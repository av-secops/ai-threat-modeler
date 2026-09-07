from io import BytesIO

import pytest
from docx import Document

from app.engine.analyzer import ThreatAnalyzer
from app.engine.parser import ArchitectureParser
from app.services.document_ingestion import (
    _extract_docx_text, _extract_pdf_text, _ocr_image, _reconstruct_pdf_lines,
)


def test_portable_ocr_is_used_when_host_tesseract_is_unavailable(monkeypatch):
    from app.services import document_ingestion

    class BrokenTesseract:
        @staticmethod
        def image_to_string(_image):
            raise RuntimeError("host executable missing")

    class FakeRapidOCR:
        def __call__(self, _image):
            return [([0, 0, 1, 1], "API to database", 0.99)], None

    monkeypatch.setattr(document_ingestion, "pytesseract", BrokenTesseract())
    monkeypatch.setattr(document_ingestion, "RapidOCR", lambda: FakeRapidOCR())
    monkeypatch.setattr(document_ingestion, "_rapid_ocr", None)

    assert _ocr_image(object()) == ("API to database", "rapidocr-onnx")


def test_wrapped_pdf_table_rows_are_reconstructed_before_classification():
    text = """Row 1: ID | Area | Known Condition
Row 2: K1 | Session | Signed evidence URLs remain valid after a user is
disabled, so prior access survives account revocation.
Row 3: K2 | GraphQL | Nested fragments have no depth or cost limit."""

    reconstructed, joined, suspicious = _reconstruct_pdf_lines(text)

    assert "user is disabled, so prior access survives account revocation" in reconstructed
    assert "Row 3: K2" in reconstructed
    assert joined == 1
    assert suspicious == 0


def _architecture_docx() -> bytes:
    document = Document()
    document.add_paragraph("3. Architecture inventory")
    table = document.add_table(rows=4, cols=7)
    rows = [
        ["ID", "Component", "Type", "Technology", "Responsibility/Data", "Trust Level", "Controls"],
        ["C1", "Customer Portal", "WebClient", "React", "customer login", "public", "WAF"],
        ["C2", "Claims API", "API", "Node.js", "claim processing", "internal", "JWT"],
        ["C3", "Claims Database", "Database", "PostgreSQL", "customer PII", "restricted", "encryption at rest"],
    ]
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            cell.text = value

    document.add_paragraph("4. Data flows")
    table = document.add_table(rows=3, cols=6)
    rows = [
        ["ID", "Source and Destination", "Protocol", "Data", "Boundary Crossing", "Evidence"],
        ["F1", "C1 -> C2", "HTTPS", "access token and claim request", "yes", "stated"],
        ["F2", "C2 -> C3", "TLS", "customer PII", "yes", "stated"],
    ]
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            cell.text = value

    document.add_paragraph("8. Known weaknesses")
    table = document.add_table(rows=2, cols=3)
    rows = [
        ["ID", "Area", "Known Condition"],
        ["K1", "Authorization", "Claims use caller-provided tenant_id without server-side ownership validation."],
    ]
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            cell.text = value
    payload = BytesIO()
    document.save(payload)
    return payload.getvalue()


def _architecture_pdf() -> bytes:
    pymupdf = pytest.importorskip("pymupdf")
    lines = [
        "3. Architecture inventory", "[Table 1]",
        "Row 1: ID | Component | Type | Technology | Responsibility/Data | Trust Level | Controls",
        "Row 2: C1 | Customer Portal | WebClient | React | customer login | public | WAF",
        "Row 3: C2 | Claims API | API | Node.js | claim processing | internal | JWT",
        "Row 4: C3 | Claims Database | Database | PostgreSQL | customer PII | restricted | encryption at rest",
        "4. Data flows", "[Table 2]",
        "Row 1: ID | Source and Destination | Protocol | Data | Boundary Crossing | Evidence",
        "Row 2: F1 | C1 -> C2 | HTTPS | access token and claim request | yes | stated",
        "Row 3: F2 | C2 -> C3 | TLS | customer PII | yes | stated",
        "8. Known weaknesses", "[Table 3]", "Row 1: ID | Area | Known Condition",
        "Row 2: K1 | Authorization | Claims use caller-provided tenant_id without server-side ownership",
        "validation.",
    ]
    document = pymupdf.open()
    page = document.new_page(width=1200, height=1000)
    page.insert_text((50, 50), "\n".join(lines), fontsize=10)
    return document.tobytes()


def test_equivalent_pdf_and_docx_produce_equivalent_model_and_findings():
    pdf_text, pdf_metadata = _extract_pdf_text(_architecture_pdf())
    docx_text, docx_metadata = _extract_docx_text(_architecture_docx())

    pdf_architecture = ArchitectureParser().parse(pdf_text)
    docx_architecture = ArchitectureParser().parse(docx_text)
    component_signature = lambda architecture: {
        (item.id, item.name, item.type) for item in architecture.components
    }
    flow_signature = lambda architecture: {
        (item.source_id, item.target_id, item.protocol.lower()) for item in architecture.flows
    }

    assert pdf_metadata["extraction_quality"] == "semantic_text_complete"
    assert docx_metadata["extraction_quality"] == "structured_text_complete"
    assert component_signature(pdf_architecture) == component_signature(docx_architecture)
    assert flow_signature(pdf_architecture) == flow_signature(docx_architecture)
    assert [item["suggested_threat_id"] for item in pdf_architecture.metadata["known_issues"]] == \
        [item["suggested_threat_id"] for item in docx_architecture.metadata["known_issues"]]

    analyzer = ThreatAnalyzer()
    pdf_result = analyzer.analyze(pdf_architecture, "PDF parity", use_local_slm=False)
    docx_result = analyzer.analyze(docx_architecture, "DOCX parity", use_local_slm=False)
    confirmed_ids = lambda result: {
        item.id.rsplit("-K", 1)[0] for item in result.threats if item.tier == "Confirmed"
    }
    assert confirmed_ids(pdf_result) == confirmed_ids(docx_result)
