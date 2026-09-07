import io
import csv
import json
import os
import re
import zipfile
import xml.etree.ElementTree as ET
from typing import Any, Dict, List, Optional, Tuple

try:
    import yaml
except Exception:  # pragma: no cover
    yaml = None

from fastapi import UploadFile

try:
    from pypdf import PdfReader
except Exception:  # pragma: no cover - optional dependency at runtime
    PdfReader = None

try:
    import pymupdf as fitz
except ImportError:  # pragma: no cover - PyMuPDF before the alias was introduced.
    try:
        import fitz
    except Exception:
        fitz = None

try:
    import pytesseract
except Exception:  # pragma: no cover - optional host adapter
    pytesseract = None

try:
    from rapidocr_onnxruntime import RapidOCR
except Exception:  # pragma: no cover - optional portable OCR backend
    RapidOCR = None

try:
    from PIL import Image
except Exception:  # pragma: no cover
    Image = None

_rapid_ocr = None

try:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
except Exception:  # pragma: no cover - optional dependency at runtime
    Document = None
    Table = Paragraph = CT_Tbl = CT_P = None


TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".rst",
    ".adoc",
    ".csv",
    ".log",
    ".json",
    ".yaml",
    ".yml",
    ".tf",
    ".hcl",
}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {".pdf", ".docx"}
MAX_DOCUMENT_BYTES = 8 * 1024 * 1024
MAX_DOCUMENTS = int(os.getenv("AEGIS_THREAT_MAX_UPLOAD_FILES", "20"))
MAX_TOTAL_BYTES = int(os.getenv("AEGIS_THREAT_MAX_UPLOAD_TOTAL_BYTES", str(32 * 1024 * 1024)))


def _read_text_bytes(raw_bytes: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("Could not decode the uploaded text document.")


_PDF_STRUCTURE_LINE = re.compile(
    r"^(?:\[?(?:page|table|paragraph)\s+\d+\]?|row\s+\d+\s*:|"
    r"(?:components?|data\s+flows?|known\s+(?:issues?|weaknesses)|trust\s+boundaries|assets?)\s*:?)$",
    re.IGNORECASE,
)


def _reconstruct_pdf_lines(text: str) -> Tuple[str, int, int]:
    """Rejoin visual line wraps without destroying document structure.

    PDF text extraction exposes the lines painted on a page, not the logical
    paragraph or table cell. A wrapped ``Row 2:`` record therefore used to be
    classified from only its first visual line. Continuation text is folded
    back into that record, while headings, list items and subsequent rows remain
    separate statements.
    """
    raw_lines = [re.sub(r"\s+", " ", line).strip() for line in (text or "").splitlines()]
    output: List[str] = []
    reconstructed = 0
    suspicious = 0

    for line in raw_lines:
        if not line:
            if output and output[-1] != "":
                output.append("")
            continue

        previous = output[-1] if output else ""
        previous_is_row = bool(re.match(r"^Row\s+\d+\s*:", previous, re.IGNORECASE))
        starts_structure = bool(
            _PDF_STRUCTURE_LINE.match(line)
            or re.match(r"^Row\s+\d+\s*:", line, re.IGNORECASE)
        )
        starts_list = bool(re.match(r"^(?:[-*]|\d+[.)])\s+", line))
        continuation = (
            bool(previous)
            and previous != ""
            and not starts_structure
            and not starts_list
            and (
                previous_is_row
                or (not re.search(r"[.!?:;]$", previous) and line[:1].islower())
            )
        )
        if continuation:
            output[-1] = f"{previous} {line}"
            reconstructed += 1
        else:
            output.append(line)

    cleaned = "\n".join(output).strip()
    for line in cleaned.splitlines():
        if re.match(r"^Row\s+\d+\s*:", line, re.IGNORECASE):
            # An authoritative row must retain at least two delimited fields.
            if line.count("|") < 1:
                suspicious += 1
        elif line and len(line) < 18 and line[-1:] not in ".:;)]" and line[:1].islower():
            suspicious += 1
    return cleaned, reconstructed, suspicious


def _ocr_image(image: Any) -> Tuple[str, str]:
    """Use host Tesseract first, then a self-contained ONNX OCR model."""
    global _rapid_ocr
    if image is None:
        return "", "unavailable"
    if pytesseract is not None:
        try:  # Import success does not imply that the host executable exists.
            text = (pytesseract.image_to_string(image) or "").strip()
            if text:
                return text, "tesseract"
        except Exception:
            pass
    if RapidOCR is not None:
        try:
            if _rapid_ocr is None:
                _rapid_ocr = RapidOCR()
            result, _ = _rapid_ocr(image)
            lines = [str(item[1]).strip() for item in result or [] if len(item) > 1 and str(item[1]).strip()]
            if lines:
                return "\n".join(lines), "rapidocr-onnx"
        except Exception:
            pass
    return "", "unavailable"


def _ocr_pdf_page(raw_bytes: bytes, page_index: int) -> str:
    """OCR one page when either optional OCR backend is available."""
    if fitz is None or Image is None:
        return ""
    try:  # pragma: no cover - exercised with mocked OCR in unit tests
        document = fitz.open(stream=raw_bytes, filetype="pdf")
        page = document.load_page(page_index)
        pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
        return _ocr_image(image)[0]
    except Exception:
        return ""


def _extract_docx_image_text(raw_bytes: bytes) -> Tuple[List[str], int, List[str]]:
    if Image is None:
        return [], 0, []
    extracted: List[str] = []
    image_count = 0
    backends = set()
    try:
        with zipfile.ZipFile(io.BytesIO(raw_bytes)) as archive:
            media_names = sorted(name for name in archive.namelist() if name.startswith("word/media/"))
            image_count = len(media_names)
            for index, media_name in enumerate(media_names, start=1):
                try:
                    image = Image.open(io.BytesIO(archive.read(media_name))).convert("RGB")
                    text, backend = _ocr_image(image)
                    if text:
                        extracted.append(f"[Embedded image {index} OCR]\n{text}")
                        backends.add(backend)
                except Exception:
                    continue
    except Exception:
        return [], 0, []
    return extracted, image_count, sorted(backends)


def _extract_pdf_text(raw_bytes: bytes) -> Tuple[str, Dict[str, str]]:
    if PdfReader is None:
        raise ValueError("PDF support is not installed on the backend.")

    reader = PdfReader(io.BytesIO(raw_bytes))
    pages = []
    image_only_pages = []
    ocr_pages = []
    reconstructed_lines = 0
    suspicious_breaks = 0
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if not text:
            text = _ocr_pdf_page(raw_bytes, index - 1)
            if text:
                ocr_pages.append(str(index))
        if text:
            text, reconstructed, suspicious = _reconstruct_pdf_lines(text)
            reconstructed_lines += reconstructed
            suspicious_breaks += suspicious
            pages.append(f"[Page {index}]\n{text}")
        else:
            image_only_pages.append(str(index))

    if not pages:
        raise ValueError("The uploaded PDF did not contain extractable text.")

    details = {
        "pages": str(len(reader.pages)),
        "image_only_pages": ",".join(image_only_pages),
        "ocr_pages": ",".join(ocr_pages),
        "reconstructed_lines": str(reconstructed_lines),
        "suspicious_line_breaks": str(suspicious_breaks),
        "extraction_quality": (
            "partial" if image_only_pages
            else "semantic_review_required" if suspicious_breaks
            else "semantic_text_complete"
        ),
        "warning": (
            f"Pages {', '.join(image_only_pages)} contained no extractable text; OCR or diagram review is required."
            if image_only_pages else ""
        ),
    }
    return "\n\n".join(pages), details


def _extract_docx_text(raw_bytes: bytes) -> Tuple[str, Dict[str, str]]:
    if Document is not None:
        document = Document(io.BytesIO(raw_bytes))
        blocks = []
        paragraph_count = 0
        table_count = 0
        # Iterating document.paragraphs followed by document.tables destroys
        # section context. Known-issue and architecture tables must remain next
        # to their headings so the architecture parser can treat them as
        # authoritative records.
        for child in document.element.body.iterchildren():
            if isinstance(child, CT_P):
                paragraph = Paragraph(child, document)
                text = paragraph.text.strip()
                if text:
                    paragraph_count += 1
                    blocks.append(f"[Paragraph {paragraph_count}] {text}")
            elif isinstance(child, CT_Tbl):
                table_count += 1
                table = Table(child, document)
                blocks.append(f"[Table {table_count}]")
                for row_index, row in enumerate(table.rows, 1):
                    values = [
                        re.sub(r"\s+", " ", cell.text).strip().replace("|", "&#124;")
                        for cell in row.cells
                    ]
                    blocks.append(f"Row {row_index}: " + " | ".join(values))

        if blocks:
            image_blocks, media_count, ocr_backends = _extract_docx_image_text(raw_bytes)
            blocks.extend(image_blocks)
            image_count = max(len(document.inline_shapes), media_count)
            unread_images = max(0, image_count - len(image_blocks))
            return "\n".join(blocks), {
                "paragraphs": str(paragraph_count),
                "tables": str(table_count),
                "embedded_images": str(image_count),
                "ocr_images": str(len(image_blocks)),
                "ocr_backends": ",".join(ocr_backends),
                "diagram_review_required": "true" if image_count else "false",
                "extraction_quality": "partial" if unread_images else "structured_text_with_image_ocr" if image_count else "structured_text_complete",
                "warning": (
                    f"{unread_images} embedded image(s) could not be read; diagram topology requires review."
                    if unread_images else "Diagram topology requires review; OCR extracts labels but does not infer connectors."
                    if image_count else ""
                ),
            }

    # Fallback parser for environments where python-docx is missing.
    try:
        namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        with zipfile.ZipFile(io.BytesIO(raw_bytes)) as archive:
            xml_payload = archive.read("word/document.xml")
        root = ET.fromstring(xml_payload)
        body = root.find("w:body", namespace)
        if body is None:
            raise ValueError("DOCX body is missing.")

        blocks = []
        paragraph_count = 0
        table_count = 0
        for child in list(body):
            local_name = child.tag.rsplit("}", 1)[-1]
            if local_name == "p":
                texts = [node.text for node in child.findall(".//w:t", namespace) if node.text]
                text = "".join(texts).strip()
                if text:
                    paragraph_count += 1
                    blocks.append(f"[Paragraph {paragraph_count}] {text}")
            elif local_name == "tbl":
                table_count += 1
                blocks.append(f"[Table {table_count}]")
                for row_index, row in enumerate(child.findall("./w:tr", namespace), 1):
                    values = []
                    for cell in row.findall("./w:tc", namespace):
                        paragraphs = []
                        for paragraph in cell.findall(".//w:p", namespace):
                            texts = [node.text for node in paragraph.findall(".//w:t", namespace) if node.text]
                            if texts:
                                paragraphs.append("".join(texts))
                        value = re.sub(r"\s+", " ", " ".join(paragraphs)).strip().replace("|", "&#124;")
                        values.append(value)
                    blocks.append(f"Row {row_index}: " + " | ".join(values))
    except Exception as exc:
        raise ValueError(f"DOCX support is not installed on the backend and fallback parsing failed: {exc}") from exc

    if not blocks:
        raise ValueError("The uploaded DOCX did not contain extractable text.")

    return "\n".join(blocks), {
        "paragraphs": str(paragraph_count), "tables": str(table_count), "embedded_images": "unknown",
        "extraction_quality": "structured_text_fallback",
        "warning": "DOCX images require separate review; paragraph and table order was preserved.",
    }


def _extract_structured_text(raw_bytes: bytes, extension: str) -> Tuple[str, Dict[str, str]]:
    raw_text = _read_text_bytes(raw_bytes).strip()
    if extension == ".json":
        parsed = json.loads(raw_text)
        flattened = _flatten_structure(parsed)
        return json.dumps(parsed, indent=2, ensure_ascii=True) + "\n\n[Structured paths]\n" + "\n".join(flattened), {
            "extraction_quality": "structured_complete", "structured_format": "json", "warning": "",
        }
    if extension in {".yaml", ".yml"} and yaml is not None:
        parsed = yaml.safe_load(raw_text)
        if _is_architecture_schema(parsed):
            rendered = _render_architecture_schema(parsed)
            return rendered, {
                "extraction_quality": "structured_complete",
                "structured_format": "yaml",
                "structured_kind": "architecture",
                "system_name": _scalar(parsed.get("system") or parsed.get("name")),
                "warning": "",
            }
        flattened = _flatten_structure(parsed)
        structured_kind = _structured_yaml_kind(parsed)
        return raw_text + "\n\n[Structured paths]\n" + "\n".join(flattened), {
            "extraction_quality": "structured_complete", "structured_format": "yaml",
            "structured_kind": structured_kind, "warning": "",
        }
    if extension == ".csv":
        rows = list(csv.reader(io.StringIO(raw_text)))
        rendered = [f"[CSV row {index}] " + " | ".join(row) for index, row in enumerate(rows, 1)]
        return "\n".join(rendered), {
            "extraction_quality": "structured_complete", "structured_format": "csv", "rows": str(len(rows)), "warning": "",
        }
    return raw_text, {"extraction_quality": "text_complete", "warning": ""}


def _structured_yaml_kind(value: Any) -> str:
    if not isinstance(value, dict):
        return "generic_configuration"
    if "services" in value and isinstance(value.get("services"), dict):
        return "docker_compose"
    if "apiVersion" in value and "kind" in value:
        return "kubernetes"
    if "jobs" in value or "stages" in value or "workflows" in value:
        return "ci_cd"
    return "generic_configuration"


def _is_architecture_schema(value: Any) -> bool:
    return (
        isinstance(value, dict)
        and isinstance(value.get("components"), list)
        and bool(value.get("components"))
        and isinstance(value.get("flows", value.get("data_flows", [])), list)
    )


def _scalar(value: Any, default: str = "") -> str:
    if value is None:
        return default
    if isinstance(value, (list, tuple, set)):
        return ", ".join(str(item) for item in value)
    if isinstance(value, dict):
        return ", ".join(f"{key}={child}" for key, child in value.items())
    return str(value)


def _render_table(rows: List[List[str]]) -> List[str]:
    rendered = []
    for index, row in enumerate(rows, 1):
        values = [re.sub(r"\s+", " ", _scalar(value)).strip().replace("|", "&#124;") for value in row]
        rendered.append(f"Row {index}: " + " | ".join(values))
    return rendered


def _render_architecture_schema(model: Dict[str, Any]) -> str:
    """Render canonical YAML as authoritative records consumed by the parser.

    The record IDs provide stable cross-references while ``Canonical ID`` keeps
    the caller's identifiers intact in the architecture IR.
    """
    components = model.get("components") or []
    flows = model.get("flows") or model.get("data_flows") or []
    by_id: Dict[str, Tuple[str, str]] = {}
    component_rows = [[
        "ID", "Canonical ID", "Component", "Type", "Technology",
        "Responsibility/Data", "Trust Level", "Controls",
    ]]
    for index, component in enumerate(components, 1):
        if not isinstance(component, dict):
            continue
        record_id = f"C{index}"
        canonical_id = _scalar(component.get("id") or component.get("component_id") or record_id)
        name = _scalar(component.get("name") or canonical_id)
        by_id[canonical_id] = (record_id, name)
        component_rows.append([
            record_id, canonical_id, name,
            _scalar(component.get("type"), "Service"),
            _scalar(component.get("technology")),
            _scalar(component.get("description") or component.get("responsibility") or component.get("data")),
            _scalar(component.get("trust") or component.get("trust_level"), "internal"),
            _scalar(component.get("controls") or component.get("properties")),
        ])

    flow_rows = [["ID", "Source and Destination", "Protocol", "Data", "Boundary Crossing", "Evidence"]]
    for index, flow in enumerate(flows, 1):
        if not isinstance(flow, dict):
            continue
        source = _scalar(flow.get("source") or flow.get("source_id"))
        target = _scalar(flow.get("target") or flow.get("target_id") or flow.get("destination"))
        source_ref = by_id.get(source, (source, source))[0]
        target_ref = by_id.get(target, (target, target))[0]
        flow_rows.append([
            f"F{index}", f"{source_ref} -> {target_ref}",
            _scalar(flow.get("protocol"), "HTTPS"),
            _scalar(flow.get("data") or flow.get("data_type"), "application data"),
            _scalar(flow.get("boundary_crossing") or flow.get("trust_boundary")),
            "Explicit structured architecture flow",
        ])

    sections = ["[Aegis Structured Architecture]", "[Table 1]", *_render_table(component_rows)]
    if len(flow_rows) > 1:
        sections.extend(["[Table 2]", *_render_table(flow_rows)])

    boundaries = model.get("trust_boundaries") or model.get("boundaries") or []
    if boundaries:
        rows = [["ID", "Boundary", "Trust Level", "Contents"]]
        for index, boundary in enumerate(boundaries, 1):
            if not isinstance(boundary, dict):
                continue
            members = [by_id.get(_scalar(item), (_scalar(item), _scalar(item)))[0]
                       for item in boundary.get("components", boundary.get("members", []))]
            rows.append([
                f"TB{index}", _scalar(boundary.get("name"), f"Boundary {index}"),
                _scalar(boundary.get("trust_level") or boundary.get("type"), "internal"),
                ", ".join(members),
            ])
        sections.extend(["[Table 3]", *_render_table(rows)])

    issues = model.get("known_issues") or model.get("known_weaknesses") or []
    if issues:
        rows = [["ID", "Area", "Known Condition"]]
        for index, issue in enumerate(issues, 1):
            if isinstance(issue, dict):
                area = _scalar(issue.get("area") or issue.get("category"), "Security")
                condition = _scalar(issue.get("condition") or issue.get("description") or issue.get("title"))
            else:
                area, condition = "Security", _scalar(issue)
            rows.append([f"K{index}", area, condition])
        sections.extend(["[Table 4]", *_render_table(rows)])

    return "\n".join(sections)


def _flatten_structure(value, path: str = "$") -> List[str]:
    lines = []
    if isinstance(value, dict):
        for key, child in value.items():
            lines.extend(_flatten_structure(child, f"{path}.{key}"))
    elif isinstance(value, list):
        for index, child in enumerate(value):
            lines.extend(_flatten_structure(child, f"{path}[{index}]"))
    else:
        lines.append(f"{path} = {value!r}")
    return lines


def _extract_text_from_bytes(filename: str, raw_bytes: bytes) -> Tuple[str, str, Dict[str, str]]:
    _, extension = os.path.splitext((filename or "").lower())
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type '{extension or 'unknown'}'. "
            "Supported formats: txt, md, rst, csv, json, yaml, yml, tf, hcl, pdf, docx."
        )

    if len(raw_bytes) > MAX_DOCUMENT_BYTES:
        raise ValueError(f"File '{filename}' exceeds the 8 MB upload limit.")

    if extension in TEXT_EXTENSIONS:
        text, details = _extract_structured_text(raw_bytes, extension)
        return text, extension, details
    if extension == ".pdf":
        text, details = _extract_pdf_text(raw_bytes)
        return text.strip(), extension, details
    if extension == ".docx":
        text, details = _extract_docx_text(raw_bytes)
        return text.strip(), extension, details

    raise ValueError(f"Unsupported file type '{extension}'.")


async def extract_documents(
    files: List[UploadFile],
    max_files: int = MAX_DOCUMENTS,
    max_total_bytes: int = MAX_TOTAL_BYTES,
) -> Tuple[str, List[Dict[str, str]]]:
    if not files:
        raise ValueError("At least one design document must be uploaded.")
    if len(files) > max_files:
        raise ValueError(
            f"{len(files)} files were submitted; this endpoint accepts at most {max_files} per analysis."
        )

    extracted_sections: List[str] = []
    metadata: List[Dict[str, str]] = []
    total_bytes = 0

    for file in files:
        filename = file.filename or "uploaded-document"
        raw_bytes = await file.read()
        # Every document is held in memory and concatenated into one description,
        # so the batch needs a ceiling of its own, not only a per-file one.
        total_bytes += len(raw_bytes)
        if total_bytes > max_total_bytes:
            raise ValueError(
                f"The submitted documents exceed the {max_total_bytes // (1024 * 1024)} MB total upload limit."
            )
        extracted_text, extension, extraction_details = _extract_text_from_bytes(filename, raw_bytes)
        if not extracted_text:
            raise ValueError(f"File '{filename}' did not contain usable text.")

        lowered = extracted_text.lower()
        role = "source_design"
        if (
            "threat modeling report" in lowered
            or "security score" in lowered and "confirmed risks" in lowered
            or "top 3 things to fix first" in lowered
        ):
            role = "reference_report"

        extracted_sections.append(
            f"Document: {filename}\n"
            f"Type: {extension.lstrip('.')}\n"
            f"Role: {role}\n"
            f"Content:\n{extracted_text}"
        )
        metadata.append(
            {
                "filename": filename,
                "type": extension.lstrip("."),
                "role": role,
                "characters": str(len(extracted_text)),
                **_source_characteristics(extracted_text),
                **extraction_details,
            }
        )

    return "\n\n---\n\n".join(extracted_sections), metadata


_TECHNOLOGY_SIGNALS = (
    "react", "next.js", "node.js", "fastapi", "graphql", "keycloak", "auth0",
    "postgresql", "mysql", "mongodb", "redis", "kafka", "rabbitmq", "s3",
    "lambda", "ec2", "rds", "eks", "azure", "gcp", "kubernetes", "docker",
    "stripe", "sendgrid", "openai", "bedrock", "mcp", "terraform",
)


def _source_characteristics(text: str) -> Dict[str, Any]:
    lowered = (text or "").lower()
    system = re.search(r"(?im)^\s*(?:system|application|project)\s*:\s*([^\n|]{2,100})", text or "")
    environment = re.search(r"(?i)\b(production|prod|staging|stage|development|dev|test)\b", text or "")
    version = re.search(r"(?im)^\s*(?:architecture\s+)?(?:version|revision|release)\s*:\s*([A-Za-z0-9._-]{1,40})", text or "")
    updated = re.search(r"(?im)^\s*(?:last\s+updated|updated|effective\s+date)\s*:\s*([^\n|]{4,40})", text or "")
    signals = [signal for signal in _TECHNOLOGY_SIGNALS if signal in lowered]
    result = {
        "technology_signals": ",".join(signals),
        "environment": environment.group(1).lower() if environment else "",
        "deployment_version": version.group(1).strip() if version else "",
        "last_updated": updated.group(1).strip() if updated else "",
    }
    if system:
        result["system_name"] = re.sub(r"\s+", " ", system.group(1)).strip()
    return result
